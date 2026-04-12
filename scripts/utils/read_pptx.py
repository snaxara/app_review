"""
Script para ler e extrair conteúdo de arquivos PowerPoint (.pptx)

Uso:
    python read_pptx.py --file caminho/arquivo.pptx
"""

import argparse
from pptx import Presentation
import sys

def read_pptx(filepath):
    """Lê um arquivo PowerPoint e extrai o conteúdo."""
    
    try:
        prs = Presentation(filepath)
        
        print("="*80)
        print(f"APRESENTAÇÃO: {filepath}")
        print("="*80)
        print(f"\nTotal de slides: {len(prs.slides)}\n")
        
        content = []
        
        for i, slide in enumerate(prs.slides, 1):
            print(f"\n{'='*80}")
            print(f"SLIDE {i}")
            print(f"{'='*80}")
            
            slide_content = {
                'slide_number': i,
                'title': '',
                'text': []
            }
            
            # Extrair texto de todas as formas no slide
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        print(f"\n{text}")
                        slide_content['text'].append(text)
                
                # Verificar se é uma tabela
                if shape.has_table:
                    print(f"\n[TABELA]")
                    for row in shape.table.rows:
                        row_data = []
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            if cell_text:
                                row_data.append(cell_text)
                        if row_data:
                            print(" | ".join(row_data))
                            slide_content['text'].append(" | ".join(row_data))
            
            content.append(slide_content)
        
        return content
        
    except ImportError:
        print("ERRO: Biblioteca python-pptx não instalada.")
        print("Instale com: pip install python-pptx")
        return None
    except Exception as e:
        print(f"ERRO ao ler arquivo: {e}")
        return None

def main():
    parser = argparse.ArgumentParser(
        description='Lê e extrai conteúdo de arquivos PowerPoint (.pptx)'
    )
    parser.add_argument(
        '--file',
        type=str,
        required=True,
        help='Caminho para o arquivo .pptx'
    )
    
    args = parser.parse_args()
    
    content = read_pptx(args.file)
    
    if content:
        print("\n" + "="*80)
        print("LEITURA CONCLUÍDA")
        print("="*80)

if __name__ == "__main__":
    main()

